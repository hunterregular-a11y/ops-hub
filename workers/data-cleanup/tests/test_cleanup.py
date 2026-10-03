import csv
import json
import tempfile
import unittest
from pathlib import Path
import sys

HERE = Path(__file__).resolve().parent
WORKER = HERE.parent
sys.path.insert(0, str(WORKER))

from cleanup import CleanupError, run  # noqa: E402

try:
    from openpyxl import Workbook, load_workbook
except ImportError:  # pragma: no cover
    Workbook = None
    load_workbook = None


CONFIG = {
    "headers": {"mode": "snake_case"},
    "normalization": {
        "trim_whitespace": True,
        "collapse_internal_whitespace": True,
        "unicode_nfkc": True,
        "blank_tokens": ["", "n/a", "na", "null"],
        "columns": {
            "email": {"normalizers": ["lowercase"]},
            "phone": {"normalizers": ["digits_only"]},
        },
    },
    "dedupe": {"enabled": True, "keys": ["email"], "case_sensitive": False},
    "validation": {
        "columns": {
            "name": {"required": True},
            "email": {"required": True, "type": "email"},
            "phone": {"regex": "^[0-9]{10,15}$"},
        }
    },
}


class CleanupTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        self.config = self.root / "config.json"
        self.config.write_text(json.dumps(CONFIG), encoding="utf-8")

    def tearDown(self):
        self.tmp.cleanup()

    def test_csv_cleanup_is_auditable_and_non_destructive(self):
        source = HERE / "fixtures" / "contacts.csv"
        original = source.read_bytes()
        outputs = run(source, self.config, self.root / "out")

        self.assertEqual(source.read_bytes(), original)
        report = json.loads(outputs["report_json"].read_text(encoding="utf-8"))
        self.assertEqual(report["metrics"]["source_rows"], 6)
        self.assertEqual(report["metrics"]["cleaned_rows"], 5)
        self.assertEqual(report["metrics"]["duplicate_rows_removed"], 1)
        self.assertEqual(report["metrics"]["exception_rows"], 3)
        self.assertEqual(report["metrics"]["validation_issue_count"], 2)
        self.assertNotIn("path", report["source"])
        self.assertEqual(report["source"]["name"], "contacts.csv")
        self.assertFalse(Path(report["outputs"]["cleaned"]).is_absolute())

        with outputs["cleaned"].open(encoding="utf-8", newline="") as handle:
            rows = list(csv.DictReader(handle))
        self.assertEqual(rows[0]["name"], "Alice Smith")
        self.assertEqual(rows[0]["email"], "alice@example.com")
        self.assertEqual(rows[0]["phone"], "7095550100")
        self.assertEqual(rows[1]["name"], "Bob Jones")

        with outputs["exceptions"].open(encoding="utf-8", newline="") as handle:
            exceptions = list(csv.DictReader(handle))
        duplicate = next(row for row in exceptions if "duplicate_of_source_row" in row["_issues"])
        self.assertEqual(duplicate["_source_row"], "3")
        self.assertIn("duplicate_of_source_row:2", duplicate["_issues"])

    @unittest.skipIf(Workbook is None, "openpyxl unavailable")
    def test_xlsx_input_produces_xlsx_cleaned_output(self):
        source = self.root / "contacts.xlsx"
        wb = Workbook()
        ws = wb.active
        ws.title = "Contacts"
        ws.append(["Name", "Email", "Phone"])
        ws.append([" Alice ", "ALICE@EXAMPLE.COM", "(709) 555-0100"])
        ws.append(["Alice", "alice@example.com", "7095550100"])
        wb.save(source)

        cfg = dict(CONFIG)
        cfg = json.loads(json.dumps(CONFIG))
        cfg["input"] = {"sheet": "Contacts"}
        cfg_path = self.root / "xlsx-config.json"
        cfg_path.write_text(json.dumps(cfg), encoding="utf-8")

        outputs = run(source, cfg_path, self.root / "out-xlsx")
        self.assertEqual(outputs["cleaned"].suffix, ".xlsx")
        wb_out = load_workbook(outputs["cleaned"], read_only=True)
        rows = list(wb_out.active.iter_rows(values_only=True))
        self.assertEqual(len(rows), 2)  # header + one deduplicated row
        self.assertEqual(rows[1][1], "alice@example.com")

    def test_refuses_to_overwrite_outputs_without_explicit_flag(self):
        source = HERE / "fixtures" / "contacts.csv"
        output_dir = self.root / "out"
        run(source, self.config, output_dir)
        with self.assertRaises(CleanupError):
            run(source, self.config, output_dir)
        outputs = run(source, self.config, output_dir, overwrite=True)
        self.assertTrue(outputs["report_json"].exists())

    def test_near_duplicate_is_not_fuzzy_merged(self):
        source = self.root / "near.csv"
        source.write_text(
            "Name,Email,Phone\nAlice,alice@example.com,7095550100\nAlice,alice+tag@example.com,7095550100\n",
            encoding="utf-8",
        )
        outputs = run(source, self.config, self.root / "near-out")
        report = json.loads(outputs["report_json"].read_text(encoding="utf-8"))
        self.assertEqual(report["metrics"]["source_rows"], 2)
        self.assertEqual(report["metrics"]["cleaned_rows"], 2)
        self.assertEqual(report["metrics"]["duplicate_rows_removed"], 0)

    def test_invalid_or_blank_dedupe_keys_are_never_auto_removed(self):
        source = self.root / "unsafe-keys.csv"
        source.write_text(
            "Name,Email,Phone\n"
            "First,not-an-email,7095550100\n"
            "Second,not-an-email,7095550101\n"
            "Blank One,,7095550102\n"
            "Blank Two,,7095550103\n",
            encoding="utf-8",
        )
        outputs = run(source, self.config, self.root / "unsafe-out")
        report = json.loads(outputs["report_json"].read_text(encoding="utf-8"))
        self.assertEqual(report["metrics"]["source_rows"], 4)
        self.assertEqual(report["metrics"]["cleaned_rows"], 4)
        self.assertEqual(report["metrics"]["duplicate_rows_removed"], 0)
        with outputs["exceptions"].open(encoding="utf-8", newline="") as handle:
            exceptions = list(csv.DictReader(handle))
        self.assertEqual(len(exceptions), 4)
        self.assertTrue(any("dedupe_skipped_invalid_key:email" in r["_issues"] for r in exceptions))
        self.assertTrue(any("dedupe_skipped_blank_key:email" in r["_issues"] for r in exceptions))

    def test_unknown_config_column_fails_closed(self):
        source = HERE / "fixtures" / "contacts.csv"
        cfg = json.loads(json.dumps(CONFIG))
        cfg["validation"]["columns"]["does_not_exist"] = {"required": True}
        cfg_path = self.root / "bad-config.json"
        cfg_path.write_text(json.dumps(cfg), encoding="utf-8")
        with self.assertRaises(CleanupError):
            run(source, cfg_path, self.root / "out")


if __name__ == "__main__":
    unittest.main()
